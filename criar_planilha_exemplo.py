# ========================================
# CRIAR PLANILHA RELATORIO_INVENTARIO.xlsx
# Cria estrutura básica para o sistema funcionar
# ========================================

import pandas as pd
import os
import sys
from datetime import datetime

def get_base_path():
    """Retorna o caminho base"""
    if getattr(sys, 'frozen', False):
        return os.path.dirname(sys.executable)
    else:
        return os.path.dirname(os.path.abspath(__file__))

DIR_BASE = get_base_path()
DIR_DATA = os.path.join(DIR_BASE, "data")
ARQUIVO_RELATORIO = os.path.join(DIR_DATA, "RELATORIO_INVENTARIO.xlsx")

print("=" * 70)
print("📊 CRIAR PLANILHA RELATORIO_INVENTARIO.xlsx")
print("=" * 70)

# Verificar se já existe
if os.path.exists(ARQUIVO_RELATORIO):
    print(f"\n⚠️ ATENÇÃO: Planilha já existe!")
    print(f"   Caminho: {ARQUIVO_RELATORIO}")
    resposta = input("\n   Deseja sobrescrever? (S/N): ").upper()
    if resposta != 'S':
        print("\n❌ Operação cancelada.")
        input("\nPressione ENTER para sair...")
        sys.exit(0)

print(f"\n📂 Criando planilha em: {ARQUIVO_RELATORIO}")

try:
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    
    # Criar workbook
    wb = Workbook()
    
    # ========================================
    # ABA CONTAGEM
    # ========================================
    ws = wb.active
    ws.title = 'CONTAGEM'
    
    # Estilos
    header_fill = PatternFill(start_color='2E75B6', end_color='2E75B6', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True, size=11)
    titulo_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
    titulo_font = Font(color='FFFFFF', bold=True, size=16)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Título
    ws.merge_cells('A1:K1')
    titulo_cell = ws['A1']
    titulo_cell.value = '📊 CONTROLE DE INVENTÁRIO - CONTAGEM'
    titulo_cell.fill = titulo_fill
    titulo_cell.font = titulo_font
    titulo_cell.alignment = Alignment(horizontal='center', vertical='center')
    
    # Linha em branco
    ws.row_dimensions[2].height = 5
    
    # Cabeçalhos (linha 4)
    headers = [
        'ITEM', 'QTD FÍSICA', 'QTD FISCAL', 'DIFERENÇA', 'STATUS',
        'ACURACIDADE %', 'EXCESSO %', 'FALTA %', 'AJUSTADO', 'COLABORADOR', 'DATA/HORA'
    ]
    
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border
    
    # Larguras das colunas
    ws.column_dimensions['A'].width = 15  # ITEM
    ws.column_dimensions['B'].width = 15  # QTD FÍSICA
    ws.column_dimensions['C'].width = 15  # QTD FISCAL
    ws.column_dimensions['D'].width = 12  # DIFERENÇA
    ws.column_dimensions['E'].width = 20  # STATUS
    ws.column_dimensions['F'].width = 15  # ACURACIDADE %
    ws.column_dimensions['G'].width = 12  # EXCESSO %
    ws.column_dimensions['H'].width = 12  # FALTA %
    ws.column_dimensions['I'].width = 12  # AJUSTADO
    ws.column_dimensions['J'].width = 20  # COLABORADOR
    ws.column_dimensions['K'].width = 20  # DATA/HORA
    
    # Dados de exemplo (opcional)
    dados_exemplo = [
        ['VK22020', 100, '', '', '', '', '', '', '', 'João Silva', ''],
        ['VK22021', 50, '', '', '', '', '', '', '', 'Maria Santos', ''],
        ['VK22022', 200, '', '', '', '', '', '', '', 'Pedro Costa', ''],
    ]
    
    for row_idx, dados in enumerate(dados_exemplo, start=5):
        for col_idx, valor in enumerate(dados, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=valor)
            cell.border = border
            cell.alignment = Alignment(horizontal='center', vertical='center')
    
    # Congelar painéis
    ws.freeze_panes = 'A5'
    
    print("   ✅ Aba CONTAGEM criada")
    
    # ========================================
    # ABA HISTORICO_AJUSTES
    # ========================================
    ws_hist = wb.create_sheet('HISTORICO_AJUSTES')
    
    # Título
    ws_hist.merge_cells('A1:J1')
    titulo_cell = ws_hist['A1']
    titulo_cell.value = '📋 HISTÓRICO DE AJUSTES'
    titulo_cell.fill = titulo_fill
    titulo_cell.font = titulo_font
    titulo_cell.alignment = Alignment(horizontal='center', vertical='center')
    
    # Linha em branco
    ws_hist.row_dimensions[2].height = 5
    
    # Cabeçalhos
    headers_hist = [
        'ITEM', 'QTD FÍSICA', 'QTD FISCAL', 'DIFERENÇA', 'STATUS',
        'ACURACIDADE %', 'EXCESSO %', 'FALTA %', 'COLABORADOR', 'DATA/HORA'
    ]
    
    for col_idx, header in enumerate(headers_hist, start=1):
        cell = ws_hist.cell(row=4, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border
    
    # Larguras das colunas
    ws_hist.column_dimensions['A'].width = 15  # ITEM
    ws_hist.column_dimensions['B'].width = 15  # QTD FÍSICA
    ws_hist.column_dimensions['C'].width = 15  # QTD FISCAL
    ws_hist.column_dimensions['D'].width = 12  # DIFERENÇA
    ws_hist.column_dimensions['E'].width = 20  # STATUS
    ws_hist.column_dimensions['F'].width = 15  # ACURACIDADE %
    ws_hist.column_dimensions['G'].width = 12  # EXCESSO %
    ws_hist.column_dimensions['H'].width = 12  # FALTA %
    ws_hist.column_dimensions['I'].width = 20  # COLABORADOR
    ws_hist.column_dimensions['J'].width = 20  # DATA/HORA
    
    # Congelar painéis
    ws_hist.freeze_panes = 'A5'
    
    print("   ✅ Aba HISTORICO_AJUSTES criada")
    
    # ========================================
    # ABA DASHBOARD (opcional)
    # ========================================
    ws_dash = wb.create_sheet('DASHBOARD')
    
    # Título
    ws_dash.merge_cells('A1:B1')
    titulo_cell = ws_dash['A1']
    titulo_cell.value = '📊 DASHBOARD DE INDICADORES'
    titulo_cell.fill = titulo_fill
    titulo_cell.font = titulo_font
    titulo_cell.alignment = Alignment(horizontal='center', vertical='center')
    
    # Linha em branco
    ws_dash.row_dimensions[2].height = 5
    
    # Cabeçalhos
    ws_dash.cell(row=4, column=1, value='Indicador').fill = header_fill
    ws_dash.cell(row=4, column=1).font = header_font
    ws_dash.cell(row=4, column=2, value='Valor').fill = header_fill
    ws_dash.cell(row=4, column=2).font = header_font
    
    # Larguras
    ws_dash.column_dimensions['A'].width = 40
    ws_dash.column_dimensions['B'].width = 25
    
    # Indicadores de exemplo
    indicadores = [
        ['📊 Total de Itens Contados', '0'],
        ['🎯 Acuracidade Geral (%)', '0%'],
        ['🔴 Qtd Itens com Excesso', '0'],
        ['⚠️ Qtd Itens com Falta', '0'],
        ['✅ Qtd Itens Exatos', '0'],
    ]
    
    for row_idx, dados in enumerate(indicadores, start=5):
        ws_dash.cell(row=row_idx, column=1, value=dados[0])
        ws_dash.cell(row=row_idx, column=2, value=dados[1])
    
    print("   ✅ Aba DASHBOARD criada")
    
    # Salvar
    wb.save(ARQUIVO_RELATORIO)
    wb.close()
    
    print(f"\n✅ SUCESSO!")
    print(f"\n📋 Planilha criada com:")
    print(f"   ✓ Aba CONTAGEM (com 3 itens de exemplo)")
    print(f"   ✓ Aba HISTORICO_AJUSTES (vazia, pronta para uso)")
    print(f"   ✓ Aba DASHBOARD (com indicadores básicos)")
    print(f"\n📂 Local: {ARQUIVO_RELATORIO}")
    print(f"📏 Tamanho: {os.path.getsize(ARQUIVO_RELATORIO)} bytes")
    
    print(f"\n💡 PRÓXIMOS PASSOS:")
    print(f"   1. Abra a planilha e adicione seus itens na aba CONTAGEM")
    print(f"   2. Preencha: ITEM, QTD FÍSICA, COLABORADOR")
    print(f"   3. Deixe as outras colunas vazias (automação preenche)")
    print(f"   4. Execute: python configurar_compartilhamento.py")
    print(f"   5. Execute a automação")

except ImportError:
    print("\n❌ ERRO: Biblioteca openpyxl não instalada!")
    print("   Execute: pip install openpyxl")
except Exception as e:
    print(f"\n❌ ERRO: {e}")
    import traceback
    traceback.print_exc()

print("\n" + "=" * 70)
input("\nPressione ENTER para sair...")
